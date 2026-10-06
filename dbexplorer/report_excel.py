from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OP_LETTER = {"INSERT": "I", "UPDATE": "U", "DELETE": "D", "MERGE": "M", "TRUNCATE": "T", "SELECT": "S",
             "REFERENCED": "?"}
OP_FILL = {"I": "C6EFCE", "U": "FFEB9C", "D": "FFC7CE", "M": "E4D1F5", "T": "FFC7CE", "S": "DDEBF7", "?": "EDEDED"}


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            sample = [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None]
            w = min(60, max(len(h) + 2, max(sample, default=8) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def write_excel(doc, path):
    names = {t["key"]: t["display"] for t in doc["tables"]}
    names.update({m["key"]: m["display"] for m in doc["modules"]})
    nm = lambda k: names.get(k, k)
    wb = Workbook()
    wb.remove(wb.active)

    edges_per_table = defaultdict(int)
    for e in doc["edges"]:
        if e["confidence"] >= 40:
            edges_per_table[e["a"]["table"]] += 1
            edges_per_table[e["b"]["table"]] += 1
    mods_per_obj = defaultdict(set)
    for u in doc["usage"]:
        mods_per_obj[u["object"]].add(u["module"])

    # Summary
    ws = wb.create_sheet("Summary")
    m = doc["meta"]
    tables = [t for t in doc["tables"] if t["kind"] == "table"]
    lines = [
        ("Server", m.get("server")), ("Database", m.get("database")),
        ("Extracted", m.get("extracted_at")), ("Analyzed", m.get("analyzed_at")),
        ("Tables", len(tables)), ("Views", len(doc["tables"]) - len(tables)),
        ("Procedures", sum(1 for x in doc["modules"] if x["type_code"] == "P")),
        ("Triggers", sum(1 for x in doc["modules"] if x["type_code"] == "TR")),
        ("Functions", sum(1 for x in doc["modules"] if x["type"] == "Function")),
        ("Relationships (confidence >= 70)", sum(1 for e in doc["edges"] if e["confidence"] >= 70)),
        ("Relationships (40-69)", sum(1 for e in doc["edges"] if 40 <= e["confidence"] < 70)),
        ("Relationships (< 40)", sum(1 for e in doc["edges"] if e["confidence"] < 40)),
        ("Declared foreign keys", sum(1 for e in doc["edges"] if "fk" in e["sources"])),
        ("Tables without primary key", sum(1 for t in tables if not t["pk"])),
        ("Objects using dynamic SQL", sum(1 for x in doc["modules"] if x["dynamic_sql"])),
        ("SQL Agent jobs touching this DB", len(doc.get("jobs", [])) if not doc["meta"].get("jobs_error")
         else "unavailable: " + doc["meta"]["jobs_error"]),
        ("Errors raised in code", sum(len(x.get("errors", [])) for x in doc["modules"])),
        ("", ""),
        ("CRUD letters", "I=Insert U=Update D=Delete M=Merge T=Truncate S=Select ?=referenced, operation unknown"),
    ]
    for row in lines:
        ws.append(row)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    for r in ws.iter_rows(max_col=1):
        r[0].font = Font(bold=True)

    _sheet(wb, "Tables",
           ["Table", "Kind", "Rows", "Prefix", "Prefix source", "Prefix coverage", "Columns", "Primary key",
            "Relationships", "Objects using it"],
           [[t["display"], t["kind"], t["rows"], t["prefix"], t["prefix_method"], t["prefix_coverage"],
             len(t["columns"]), ", ".join(t["pk"]) or "(none)", edges_per_table[t["key"]],
             len(mods_per_obj[t["key"]])] for t in doc["tables"]])

    def values_text(c):
        p = c.get("profile") or {}
        if c.get("sensitive") or not p.get("values"):
            return ""
        total = sum(n for _, n in p["values"]) or 1
        return ", ".join(f"{'NULL' if v is None else v} ({100 * n / total:.0f}%)" for v, n in p["values"][:15])

    def prof(c, k):
        return (c.get("profile") or {}).get(k)

    _sheet(wb, "Columns", ["Table", "Column", "Logical name", "Type", "PK", "Nullable", "Identity", "Sensitive",
                           "CHECK", "Default", "Null %", "Distinct", "Min", "Max", "Values", "Formats", "Examples"],
           [[t["display"], c["name"], c["logical"], c["type"], "Y" if c["pk"] else "", "Y" if c["nullable"] else "",
             "Y" if c["identity"] else "", c.get("sensitive", ""), c.get("check", ""), c.get("default", ""),
             prof(c, "null_pct"), prof(c, "distinct"), str(prof(c, "min") or ""), str(prof(c, "max") or ""),
             values_text(c), " / ".join(f[0] for f in prof(c, "formats") or []),
             "" if c.get("sensitive") else ", ".join(map(str, prof(c, "examples") or []))]
            for t in doc["tables"] for c in t["columns"]], {"Values": 50, "CHECK": 40})

    _sheet(wb, "Test Considerations", ["Table", "Consideration"],
           [[t["display"], n] for t in doc["tables"] for n in t.get("test_notes", [])], {"Consideration": 140})

    procs = sorted([m for m in doc["modules"] if m.get("risk")], key=lambda m: -m["risk"]["score"])
    _sheet(wb, "Procedure Risk", ["Object", "Type", "Risk score", "Why", "Tables written", "Fires triggers",
                                  "Parameters", "Code lines", "Branches", "Cursors", "TRY/CATCH", "Modified"],
           [[m["display"], m["type"], m["risk"]["score"], ", ".join(m["risk"]["reasons"]),
             ", ".join(nm(k) for k in m.get("writes_tables", [])), ", ".join(nm(k) for k in m.get("fires_triggers", [])),
             ", ".join(p["name"] + ("" if p["default"] is None else f" = {p['default']}") for p in m.get("parameters", [])),
             m["metrics"].get("lines"), m["metrics"].get("branches"), m["metrics"].get("cursors"),
             "yes" if m["metrics"].get("try_catch") else "no", (m.get("modified") or "")[:10]] for m in procs],
           {"Why": 50, "Parameters": 50})

    rel_rows = []
    for e in doc["edges"]:
        child, parent = (e["b"], e["a"]) if e["parent"] == "a" else (e["a"], e["b"])
        s = e["sources"]
        d = s.get("data", {})
        rel_rows.append([
            nm(child["table"]), child["column"], nm(parent["table"]), parent["column"],
            "many-to-one" if e["parent"] else "undetermined", e["confidence"],
            s.get("fk", {}).get("name", ""), s.get("code", {}).get("count", ""),
            ", ".join(list(s.get("code", {}).get("modules", {}))[:5]),
            s.get("name", {}).get("logical", ""),
            "" if not d or d.get("error") else f"{100 - d['orphan_pct']:.1f}% of {d['sampled']}",
            "; ".join(e["evidence"]),
        ])
    ws = _sheet(wb, "Relationships",
                ["From table (child)", "From column", "To table (parent)", "To column", "Cardinality", "Confidence",
                 "Declared FK", "Code joins", "Joined in", "Logical name", "Data match", "Evidence"],
                rel_rows, {"Evidence": 70, "Joined in": 50})
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        v = row[0].value or 0
        row[0].fill = PatternFill("solid", fgColor="C6EFCE" if v >= 70 else "FFEB9C" if v >= 40 else "FFC7CE")

    # CRUD matrix: tables x procedures/triggers/functions/views
    letters = defaultdict(set)
    for u in doc["usage"]:
        letters[(u["object"], u["module"])].add(OP_LETTER[u["op"]])
    used_objs = sorted({u["object"] for u in doc["usage"]} & {t["key"] for t in doc["tables"]}, key=nm)
    mods = sorted({u["module"] for u in doc["usage"]}, key=lambda k: (k not in {
        x["key"] for x in doc["modules"] if x["type_code"] == "P"}, nm(k)))
    ws = wb.create_sheet("CRUD Matrix")
    ws.append(["Table \\ Object"] + [nm(k) for k in mods])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(text_rotation=90, vertical="bottom", horizontal="center")
    ws.row_dimensions[1].height = 180
    order = "IUDMTS?"
    for obj in used_objs:
        row = [nm(obj)]
        for mk in mods:
            ls = letters.get((obj, mk))
            row.append("".join(ch for ch in order if ch in ls) if ls else "")
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=2):
        for c in row:
            if c.value:
                c.fill = PatternFill("solid", fgColor=OP_FILL[c.value[0]])
                c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 36
    for i in range(2, len(mods) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 4.5
    ws.freeze_panes = "B2"

    mtype = {x["key"]: x["type"] for x in doc["modules"]}
    _sheet(wb, "Table Usage",
           ["Table/View", "Object", "Object type", "Operation", "Columns written", "Columns read", "Lines",
            "Dynamic SQL", "Notes"],
           [[nm(u["object"]), nm(u["module"]), mtype.get(u["module"], ""), u["op"], ", ".join(u["cols_written"]),
             ", ".join(u["cols_read"]), ", ".join(map(str, u["lines"])), "Y" if u["dynamic"] else "",
             "; ".join(u["details"])] for u in doc["usage"]],
           {"Columns read": 50, "Columns written": 40})

    _sheet(wb, "Error Catalog", ["Object", "Type", "Line", "Kind", "Number", "Severity", "Message", "Raised when",
                                 "Inside CATCH", "Warning"],
           [[m["display"], m["type"], e["line"], e["kind"], e["number"], e["severity"], e["message"],
             e["condition"], "Y" if e["in_catch"] else "", e.get("warning", "")]
            for m in doc["modules"] for e in m.get("errors", [])], {"Message": 50, "Raised when": 60})

    jobs = doc.get("jobs", [])
    _sheet(wb, "Jobs", ["Job", "Enabled", "Schedule", "Last run", "Last status", "Runs 30d", "Failures 30d",
                        "Avg duration", "Max duration", "Procedures (incl. nested)", "Tables written", "Triggers fired",
                        "Risk notes"],
           [[j["name"], "Y" if j["enabled"] else "N",
             "; ".join(s["text"] + ("" if s["enabled"] else " (disabled)") for s in j["schedules"]) or "none",
             j["history"].get("last_run"), j["history"].get("last_status"), j["history"].get("runs_30d"),
             j["history"].get("failures_30d"), j["history"].get("avg"), j["history"].get("max"),
             ", ".join(nm(k) for k in j["procedures"]), ", ".join(nm(k) for k in j["writes"]),
             ", ".join(nm(k) for k in j["fires_triggers"]), " | ".join(j["warnings"])] for j in jobs]
           or [[doc["meta"].get("jobs_error") or "No SQL Agent jobs touch this database."]],
           {"Risk notes": 80, "Schedule": 40})
    _sheet(wb, "Job Steps", ["Job", "Step", "Name", "Subsystem", "Database", "On success", "On failure", "Retries",
                             "Runs", "Command"],
           [[j["name"], st["id"], st["name"], st["subsystem"], st["database"], st["on_success"], st["on_fail"],
             st["retries"], ", ".join(nm(c) for c in st["calls"]), st["command"][:500]]
            for j in jobs for st in j["steps"]], {"Command": 80})

    _sheet(wb, "Calls", ["Caller", "Caller type", "Calls", "Lines"],
           [[m["display"], m["type"], nm(c["target"]), ", ".join(map(str, c["lines"]))]
            for m in doc["modules"] for c in m["calls"]])

    review = []
    for m in doc["modules"]:
        if m["encrypted"]:
            review.append([m["display"], m["type"], "Encrypted - definition not readable", ""])
        if m["dynamic_sql"]:
            review.append([m["display"], m["type"], "Dynamic SQL - tables built at runtime may be missing",
                           "lines " + ", ".join(map(str, m["dynamic_lines"]))])
        if m["unresolved"]:
            review.append([m["display"], m["type"], "Names not found in this database", ", ".join(m["unresolved"])])
        if m["external"]:
            review.append([m["display"], m["type"], "Cross-database / linked server references",
                           ", ".join(m["external"])])
    for err in doc["errors"]:
        review.append([err["object"], "", "Analyzer error", err["error"]])
    _sheet(wb, "Review Needed", ["Object", "Type", "Issue", "Detail"], review, {"Detail": 80})

    _sheet(wb, "Logical Keys", ["Logical name", "Tables", "Status", "Columns"],
           [[g["logical"], g["tables"], g["status"], ", ".join(g.get("members", []))] for g in doc["key_groups"]],
           {"Columns": 100})

    wb.save(path)
